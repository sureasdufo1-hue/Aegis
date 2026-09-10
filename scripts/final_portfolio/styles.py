#!/usr/bin/env python3
"""
Monochrome styling helpers and document formatting functions for Aegis SOC Master Portfolio.
Adheres strictly to epoko77-ai/im-not-ai guidelines, pure black/grayscale palette, and high-readability Korean.
"""

from pathlib import Path
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

BASE_DIR = Path(__file__).parent.parent.parent.resolve()
IMAGE_DIR_COLOR = BASE_DIR / "docs" / "ai" / "evidence_annotated"


def set_cell_background(cell, fill_hex):
    """Set background color of a table cell (monochrome tint)."""
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith("shd"):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    """Set cell padding in dxa (20 dxa = 1 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table, color="B0B0B0", sz="4", val="single"):
    """Apply clean neutral gray borders to the table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def format_run(run, font_name="맑은 고딕", size_pt=10, bold=False, color_rgb=(35, 35, 35)):
    """Apply font styling to a text run using strictly black/gray RGB values."""
    run.font.name = font_name
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)
    rFonts.set(qn("w:eastAsia"), font_name)
    rFonts.set(qn("w:cs"), font_name)
    if size_pt:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    if color_rgb:
        run.font.color.rgb = RGBColor(*color_rgb)


def add_heading_1(doc, text):
    """Heading 1: Pure Black, 16pt Bold."""
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    format_run(run, font_name="맑은 고딕", size_pt=16, bold=True, color_rgb=(0, 0, 0))
    return p


def add_heading_2(doc, text):
    """Heading 2: Deep Charcoal Black, 13pt Bold."""
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.space_before = Pt(15)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    format_run(run, font_name="맑은 고딕", size_pt=13, bold=True, color_rgb=(20, 20, 20))
    return p


def add_heading_3(doc, text):
    """Heading 3: Dark Charcoal, 11pt Bold."""
    p = doc.add_paragraph(style="Heading 3")
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    format_run(run, font_name="맑은 고딕", size_pt=11, bold=True, color_rgb=(40, 40, 40))
    return p


def add_heading_4(doc, text):
    """Heading 4: Charcoal, 10pt Bold."""
    p = doc.add_paragraph(style="Heading 4")
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    format_run(run, font_name="맑은 고딕", size_pt=10, bold=True, color_rgb=(50, 50, 50))
    return p


def add_body_p(doc, text, bold_prefix=None, space_after=4):
    """Standard body paragraph in neutral dark gray with comfortable line spacing."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        format_run(r_pre, font_name="맑은 고딕", size_pt=10, bold=True, color_rgb=(0, 0, 0))
    r = p.add_run(text)
    format_run(r, font_name="맑은 고딕", size_pt=10, bold=False, color_rgb=(35, 35, 35))
    return p


def add_bullet_p(doc, text, bold_prefix=None, space_after=3):
    """Bulleted list item with pure black bullet."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.line_spacing = 1.2
    r_bullet = p.add_run("• ")
    format_run(r_bullet, font_name="맑은 고딕", size_pt=10, bold=True, color_rgb=(0, 0, 0))
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        format_run(r_pre, font_name="맑은 고딕", size_pt=10, bold=True, color_rgb=(0, 0, 0))
    r = p.add_run(text)
    format_run(r, font_name="맑은 고딕", size_pt=10, bold=False, color_rgb=(35, 35, 35))
    return p


def add_callout_box(doc, title_text, body_text, accent_color="262626", bg_color="F5F5F5"):
    """Monochrome callout box with thick charcoal/black left border and light gray tint background."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=100, bottom=100, left=160, right=160)

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:color="D0D0D0"/>'
        f'  <w:left w:val="single" w:sz="36" w:color="{accent_color}"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:color="D0D0D0"/>'
        f'  <w:right w:val="single" w:sz="4" w:color="D0D0D0"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    if title_text:
        r_t = p.add_run(title_text + "\n")
        format_run(r_t, font_name="맑은 고딕", size_pt=10, bold=True, color_rgb=(0, 0, 0))
    r_b = p.add_run(body_text)
    format_run(r_b, font_name="맑은 고딕", size_pt=9.5, bold=False, color_rgb=(40, 40, 40))
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_code_box(doc, code_text):
    """Monochrome code box with dark charcoal left border."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    set_cell_background(cell, "EFEFEF")
    set_cell_margins(cell, top=90, bottom=90, left=140, right=140)

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:color="D0D0D0"/>'
        f'  <w:left w:val="single" w:sz="24" w:color="262626"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:color="D0D0D0"/>'
        f'  <w:right w:val="single" w:sz="4" w:color="D0D0D0"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(code_text.strip())
    format_run(r, font_name="Consolas", size_pt=8.5, bold=False, color_rgb=(20, 20, 20))
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_custom_table(doc, headers, rows_data, dark_header=True):
    """Add a structured monochrome table with Dark Charcoal headers and alternating gray shading."""
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="B0B0B0")

    # Header row: Dark Charcoal background with White bold text
    for c_idx, h_text in enumerate(headers):
        cell = table.rows[0].cells[c_idx]
        cell.text = h_text
        set_cell_background(cell, "262626" if dark_header else "E0E0E0")
        set_cell_margins(cell, top=80, bottom=80, left=90, right=90)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_run(
            p.runs[0],
            font_name="맑은 고딕",
            size_pt=9,
            bold=True,
            color_rgb=(255, 255, 255) if dark_header else (0, 0, 0),
        )

    # Data rows: Alternating White and Light Gray shading
    for r_idx, r_data in enumerate(rows_data, start=1):
        bg = "F2F2F2" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = str(val)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
            p = cell.paragraphs[0]
            val_str = str(val)
            if (
                val_str
                in (
                    "PASS",
                    "FAIL",
                    "BLOCKED",
                    "CRITICAL",
                    "HIGH",
                    "MEDIUM",
                    "LOW",
                    "P1",
                    "P2",
                    "P3",
                    "P4",
                    "통과",
                    "개선",
                    "유지",
                    "Tier 4 (Adaptive)",
                    "Tier 3 (Repeatable)",
                )
                or len(val_str) <= 10
            ):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            is_bold = (
                val_str in ("PASS", "CRITICAL", "HIGH", "P1", "FAIL", "통과")
                or "92.31%" in val_str
                or "0.00%" in val_str
                or "85.71%" in val_str
            )
            color = (0, 0, 0) if is_bold else (40, 40, 40)
            format_run(p.runs[0], font_name="맑은 고딕", size_pt=8.5, bold=is_bold, color_rgb=color)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


def add_evidence_figure(doc, img_name_or_path, caption_title, desc_text, meta_data=None):
    """High-res ORIGINAL COLOR annotated screenshot with monochrome caption and 4-column metadata table."""
    if isinstance(img_name_or_path, (str, Path)):
        img_path = Path(img_name_or_path)
        if not img_path.is_absolute():
            img_path = IMAGE_DIR_COLOR / img_name_or_path
    else:
        img_path = IMAGE_DIR_COLOR / str(img_name_or_path)

    if not img_path.exists():
        print(f"Warning: Screenshot not found: {img_path}")
        return

    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture(str(img_path), width=Inches(6.25))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(4)
    r_cap = p_cap.add_run(caption_title)
    format_run(r_cap, font_name="맑은 고딕", size_pt=9.5, bold=True, color_rgb=(0, 0, 0))

    if desc_text:
        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(2)
        p_desc.paragraph_format.space_after = Pt(4)
        p_desc.paragraph_format.line_spacing = 1.15
        r_desc = p_desc.add_run(desc_text)
        format_run(r_desc, font_name="맑은 고딕", size_pt=9, bold=False, color_rgb=(50, 50, 50))

    if meta_data:
        tbl_meta = doc.add_table(rows=3, cols=4)
        tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl_meta, color="B0B0B0")

        # Row 0: 증적 번호 | ID | 검증 결과 | RESULT
        r0 = tbl_meta.rows[0].cells
        r0[0].text = "증적 번호"
        r0[1].text = meta_data.get("id", "-")
        r0[2].text = "검증 결과"
        r0[3].text = meta_data.get("result", "정상 (PASS)")
        for idx in [0, 2]:
            set_cell_background(r0[idx], "262626")
            set_cell_margins(r0[idx], top=60, bottom=60, left=80, right=80)
            p = r0[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            format_run(p.runs[0], font_name="맑은 고딕", size_pt=8.5, bold=True, color_rgb=(255, 255, 255))
        for idx in [1, 3]:
            set_cell_background(r0[idx], "FFFFFF")
            set_cell_margins(r0[idx], top=60, bottom=60, left=80, right=80)
            p = r0[idx].paragraphs[0]
            if idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            format_run(
                p.runs[0],
                font_name="맑은 고딕",
                size_pt=8.5,
                bold=(idx == 3),
                color_rgb=(0, 0, 0) if idx == 3 else (40, 40, 40),
            )

        # Row 1: 증적 명칭 | TITLE | 점검 대상 | TARGET
        r1 = tbl_meta.rows[1].cells
        r1[0].text = "증적 명칭"
        r1[1].text = meta_data.get("title", "-")
        r1[2].text = "점검 대상"
        r1[3].text = meta_data.get("target", "-")
        for idx in [0, 2]:
            set_cell_background(r1[idx], "262626")
            set_cell_margins(r1[idx], top=60, bottom=60, left=80, right=80)
            p = r1[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            format_run(p.runs[0], font_name="맑은 고딕", size_pt=8.5, bold=True, color_rgb=(255, 255, 255))
        for idx in [1, 3]:
            set_cell_background(r1[idx], "FFFFFF")
            set_cell_margins(r1[idx], top=60, bottom=60, left=80, right=80)
            p = r1[idx].paragraphs[0]
            format_run(p.runs[0], font_name="맑은 고딕", size_pt=8.5, bold=False, color_rgb=(40, 40, 40))

        # Row 2 (merged): 핵심 검증 내용 | DETAILS
        r2 = tbl_meta.rows[2].cells
        r2[0].text = "핵심 검증 내용"
        set_cell_background(r2[0], "262626")
        set_cell_margins(r2[0], top=60, bottom=60, left=80, right=80)
        p0 = r2[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_run(p0.runs[0], font_name="맑은 고딕", size_pt=8.5, bold=True, color_rgb=(255, 255, 255))

        merged_c = r2[1].merge(r2[3])
        merged_c.text = meta_data.get("details", "-")
        set_cell_background(merged_c, "FFFFFF")
        set_cell_margins(merged_c, top=60, bottom=60, left=80, right=80)
        p_m = merged_c.paragraphs[0]
        p_m.paragraph_format.line_spacing = 1.15
        format_run(p_m.runs[0], font_name="맑은 고딕", size_pt=8.5, bold=False, color_rgb=(40, 40, 40))

    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def add_toc_field(doc):
    """Inserts native Microsoft Word TOC field XML."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run()

    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls("w"))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> TOC \o "1-3" \h \z \u </w:instrText>' % nsdecls("w"))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls("w"))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls("w"))

    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)


def enable_update_fields(doc):
    """Instruct Word to refresh fields (including TOC) on open."""
    settings = doc.settings.element
    updateFields = parse_xml(r'<w:updateFields %s w:val="true"/>' % nsdecls("w"))
    settings.append(updateFields)
